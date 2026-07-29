import logging
import os
import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from src.core.models import BusinessItem
from src.core.interfaces.exporter import ExporterInterface

logger = logging.getLogger(__name__)

class ExcelExporter(ExporterInterface):
    """Concrete openpyxl implementation of ExporterInterface."""
    
    def export(self, items: List[BusinessItem], file_path: str = None) -> str:
        """
        Exports a list of BusinessItem models to a local Excel sheet.
        Injects today's date in the filename and returns the final saved path.
        """
        today_str = datetime.date.today().strftime("%Y-%m-%d")
        
        if not file_path:
            file_path = f"leads_{today_str}.xlsx"
        elif os.path.isdir(file_path):
            file_path = os.path.join(file_path, f"leads_{today_str}.xlsx")
        else:
            # Ensure file path has today's date in filename
            dir_name, file_name = os.path.split(file_path)
            base, ext = os.path.splitext(file_name)
            if today_str not in base:
                separator = "" if base.endswith("_") or base.endswith("-") else "_"
                file_name = f"{base}{separator}{today_str}{ext}"
                file_path = os.path.join(dir_name, file_name)

        logger.info(f"Exporting {len(items)} items to {file_path}...")
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Scraped Leads"
            
            # Setup columns and headers
            headers = [
                "Business Name", "Address", "Phone", "Website", 
                "Rating", "Reviews", "Google Maps URL"
            ]
            
            # Write header
            ws.append(headers)
            
            # Apply styling to headers
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # Write data rows
            for item in items:
                ws.append([
                    item.name,
                    item.address,
                    item.phone,
                    item.website,
                    item.rating,
                    item.reviews_count,
                    item.maps_url
                ])
            
            # Style data cells and auto-fit columns
            for row in range(2, len(items) + 2):
                for col in range(1, len(headers) + 1):
                    cell = ws.cell(row=row, column=col)
                    cell.font = Font(name="Calibri", size=11)
                    cell.border = thin_border
                    # Align numeric fields (Rating: col 5, Reviews: col 6) to the right
                    if col in (5, 6):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")
            
            # Set header height
            ws.row_dimensions[1].height = 25
            
            # Auto-fit columns with safety margin
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
            # Ensure parent directory exists before saving
            dir_name = os.path.dirname(file_path)
            if dir_name:
                os.makedirs(dir_name, exist_ok=True)

            wb.save(file_path)
            logger.info("Excel export completed successfully.")
            return file_path
            
        except Exception as e:
            logger.error(f"Failed to export to Excel: {e}")
            raise IOError(f"Failed to write Excel file: {e}") from e
