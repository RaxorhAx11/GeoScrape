import unittest
import os
import tempfile
import datetime
from openpyxl import load_workbook
from src.core.models import BusinessItem
from src.exporter.excel_exporter import ExcelExporter

class TestExcelExporter(unittest.TestCase):
    """Unit tests for the ExcelExporter class."""
    
    def setUp(self):
        self.exporter = ExcelExporter()
        self.temp_dir = tempfile.TemporaryDirectory()
        self.test_file_path = os.path.join(self.temp_dir.name, "test_leads.xlsx")
        
        self.test_items = [
            BusinessItem(
                name="Test Cafe",
                rating=4.5,
                reviews_count=120,
                category="Cafe",
                address="123 Main St, Austin, TX",
                phone="+1 512-555-0199",
                website="https://testcafe.com",
                maps_url="https://google.com/maps/place/1",
                email="hello@testcafe.com",
                linkedin="https://linkedin.com/company/testcafe",
                facebook="https://facebook.com/testcafe",
                instagram="https://instagram.com/testcafe"
            ),
            BusinessItem(
                name="Dental Clinic",
                rating=4.8,
                reviews_count=85,
                category="Dentist",
                address="456 Medical Pkwy, Austin, TX",
                phone="+1 512-555-0288",
                website="N/A",
                maps_url="https://google.com/maps/place/2"
            )
        ]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_export_structure(self):
        # Export the items
        saved_path = self.exporter.export(self.test_items, self.test_file_path)
        
        # Verify the file is created
        self.assertTrue(os.path.exists(saved_path))
        
        # Verify that today's date was injected in the file name
        today_str = datetime.date.today().strftime("%Y-%m-%d")
        self.assertIn(today_str, saved_path)
        
        # Load and verify content structure
        wb = load_workbook(saved_path)
        self.assertIn("Scraped Leads", wb.sheetnames)
        
        ws = wb["Scraped Leads"]
        
        # Verify headers
        headers = [cell.value for cell in ws[1]]
        expected_headers = [
            "Business Name", "Address", "Phone", "Website", 
            "Email", "LinkedIn", "Facebook", "Instagram",
            "Rating", "Reviews", "Google Maps URL"
        ]
        self.assertEqual(headers, expected_headers)
        
        # Verify rows
        self.assertEqual(ws.max_row, 3) # 1 header row + 2 data rows
        
        # Check first data row values with enrichment data
        row2 = [cell.value for cell in ws[2]]
        self.assertEqual(row2[0], "Test Cafe")
        self.assertEqual(row2[1], "123 Main St, Austin, TX")
        self.assertEqual(row2[2], "+1 512-555-0199")
        self.assertEqual(row2[3], "https://testcafe.com")
        self.assertEqual(row2[4], "hello@testcafe.com")
        self.assertEqual(row2[5], "https://linkedin.com/company/testcafe")
        self.assertEqual(row2[6], "https://facebook.com/testcafe")
        self.assertEqual(row2[7], "https://instagram.com/testcafe")
        self.assertEqual(row2[8], 4.5)
        self.assertEqual(row2[9], 120)
        self.assertEqual(row2[10], "https://google.com/maps/place/1")

        # Check second data row values with default N/A fields
        row3 = [cell.value for cell in ws[3]]
        self.assertEqual(row3[0], "Dental Clinic")
        self.assertEqual(row3[4], "N/A")
        self.assertEqual(row3[5], "N/A")
        self.assertEqual(row3[6], "N/A")
        self.assertEqual(row3[7], "N/A")
